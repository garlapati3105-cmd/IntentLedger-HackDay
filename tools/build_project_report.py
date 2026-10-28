from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "IntentLedger_Project_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep(paragraph, next_paragraph=False):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext" if next_paragraph else "w:keepLines")
    p_pr.append(keep)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, "26345C")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[i])
    for row_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2:
                set_cell_shading(cell, "F3F6FC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            r.font.size = Pt(8.7)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
for style_name, size, before, after in (("Title", 27, 0, 12), ("Heading 1", 18, 16, 7), ("Heading 2", 13, 11, 5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph(style="Title")
title.add_run("IntentLedger Project Report")
title_pr = title._p.get_or_add_pPr()
title_border = title_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_pr.remove(title_border)
# Override any border inherited from Word's built-in Title style.
empty_border = OxmlElement("w:pBdr")
for edge_name in ("top", "left", "bottom", "right", "between", "bar"):
    edge = OxmlElement(f"w:{edge_name}")
    edge.set(qn("w:val"), "nil")
    empty_border.append(edge)
title_pr.append(empty_border)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = subtitle.add_run("AI intent tracking benchmark and evaluation workspace")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(45, 58, 92)
meta = doc.add_paragraph()
meta.add_run("Prepared for hackathon demonstration and technical review\n").italic = True
meta.add_run("Project stack: Python, Streamlit, Google Gen AI SDK, hosted Gemma model")

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "IntentLedger is a small, transparent evaluation application that tests whether an AI model can reconstruct a user's final intended state after a multi-turn conversation changes direction. It focuses on five common instruction-tracking failures: accepted corrections, rejected suggestions, partial updates, cancellations, and unresolved ambiguity. The application compares a simple baseline extraction prompt with an improved prompt that explicitly defines how those situations should be handled."
)
doc.add_paragraph(
    "The project is designed for a live hackathon demonstration rather than as a statistically representative benchmark. Its main strength is explainability: every run exposes the conversation, expected state, predicted state, field-level comparison, normalized match rate, and PASS, FAIL, or ERROR classification. Results are generated only from real application state; the interface does not fabricate scores or charts."
)

doc.add_heading("Project Objectives", level=1)
add_bullets(doc, [
    "Measure whether a model preserves the user's latest confirmed intent across conversation turns.",
    "Compare a naive extraction prompt with an instruction-aware prompt on the same cases and expected states.",
    "Make failures easy to diagnose through field-level expected-versus-predicted comparisons.",
    "Separate model or configuration failures from evaluation failures using a distinct ERROR state.",
    "Present the benchmark in a professional, judge-friendly Streamlit interface suitable for a live demonstration.",
])

doc.add_heading("Problem Statement", level=1)
doc.add_paragraph(
    "Conversational assistants often treat each new turn as additional information without correctly resolving its relationship to earlier turns. A corrected meeting time may coexist with the original time, an assistant suggestion may be mistaken for user intent, or a speculative value may be treated as confirmed. These errors are difficult to detect when an application shows only a final response. IntentLedger converts the conversation into a structured state and compares that state against human-authored ground truth so the precise failure is visible."
)

doc.add_heading("System Architecture", level=1)
doc.add_paragraph(
    "The application keeps presentation, model interaction, prompt construction, normalization, case loading, and scoring in separate modules. This separation allows the interface to change without altering evaluation semantics."
)
add_table(doc,
    ["Layer", "Primary file", "Responsibility"],
    [
        ("Presentation", "app.py and ui.py", "Streamlit state, controls, navigation, result rendering, responsive styling, and API status display."),
        ("Case loading", "engine/loader.py", "Loads the five JSON benchmark cases from tests/cases in deterministic filename order."),
        ("Prompting", "engine/prompts.py", "Builds the baseline or improved system prompt and formats the transcript."),
        ("Model access", "engine/extractor.py", "Reads configuration, calls the Google Gen AI SDK, parses JSON, and classifies model errors."),
        ("Normalization", "engine/normalizer.py", "Canonicalizes times, currency, numbers, lists, and case-insensitive strings."),
        ("Scoring", "engine/scorer.py", "Produces per-field outcomes, match rate, clarification behavior, and PASS, FAIL, or ERROR verdicts."),
        ("Verification", "tests", "Covers configuration, parsing, normalization, scoring, clarification handling, and summaries."),
    ], [1.05, 1.55, 4.55]
)

doc.add_heading("Evaluation Data Flow", level=1)
add_numbered(doc, [
    "The user selects one of the five benchmark scenarios or starts a run across all cases.",
    "The application retrieves the conversation and human-authored expected state from the case JSON.",
    "The prompt builder selects the baseline or improved system prompt and formats the conversation transcript.",
    "The extractor calls the configured model through the Google Gen AI SDK and requests deterministic JSON output.",
    "The response parser removes optional code fences and extracts the first valid JSON object.",
    "The scorer normalizes comparable values and evaluates every expected field as matched, mismatched, or missing; extra predicted fields are marked unexpected.",
    "The UI stores the ScoreResult in Streamlit session state and displays the verdict, metrics, field differences, raw output, and troubleshooting guidance when required.",
])

doc.add_heading("Benchmark Scenarios", level=1)
doc.add_paragraph("The benchmark contains five hand-crafted cases. Their meaning and expected states are preserved by the redesign.")
add_table(doc,
    ["Case", "Scenario", "Core challenge", "Critical expected behavior"],
    [
        ("01", "Accepted Correction", "Meeting time changes from 2 PM to 3 PM.", "Keep the latest time and preserve the location and attendees."),
        ("02", "Rejected Suggestion", "The assistant proposes La Maison and the user rejects it.", "Retain Bella Italia and exclude the rejected suggestion."),
        ("03", "Partial Update", "The trip budget changes from 2000 to 3000.", "Update only the budget while retaining destination, duration, and accommodation."),
        ("04", "Cancellation", "A newsletter subscription is later cancelled.", "Represent cancellation as the final action with cancelled status."),
        ("05", "Ambiguous Instruction", "The user requests a warm color but confirms no specific palette.", "Mark the color as ambiguous and request clarification instead of confirming a guess."),
    ], [0.55, 1.25, 2.5, 2.85]
)

doc.add_heading("Evaluation Modes", level=1)
doc.add_heading("Baseline Prompt", level=2)
doc.add_paragraph(
    "The baseline prompt asks the model to extract the final intended state as JSON. It does not give dedicated rules for corrections, rejected suggestions, partial updates, cancellations, or ambiguity. This provides a simple comparison point for the experiment."
)
doc.add_heading("Improved Prompt", level=2)
doc.add_paragraph(
    "The improved prompt defines five explicit instruction-tracking rules. It requires the latest correction to replace the old value, rejected assistant suggestions to be excluded, unchanged fields to survive partial updates, cancellations to become the active final intent, and unresolved details to be marked AMBIGUOUS with clarification metadata."
)

doc.add_heading("Scoring and Verdict Logic", level=1)
add_table(doc,
    ["Outcome", "Meaning", "Effect"],
    [
        ("Match", "Normalized predicted value equals the normalized expected value.", "Counts toward field-match rate."),
        ("Mismatch", "The field exists but contains a different value.", "Causes FAIL."),
        ("Missing", "An expected field is absent from the prediction.", "Causes FAIL."),
        ("Unexpected", "The prediction contains a field not present in expected state.", "Displayed for inspection but does not fail by itself."),
        ("ERROR", "No valid prediction is available because the model call, configuration, dependency, or JSON parse failed.", "The case is not scored and retains its original error classification."),
    ], [1.0, 4.05, 2.1]
)
doc.add_paragraph(
    "Field-match rate equals matched expected fields divided by total expected fields. ERROR results are excluded from the aggregate average. The ambiguity case also checks whether clarification_needed is true; failing that requirement produces FAIL even if other fields match."
)

doc.add_heading("Normalization Rules", level=1)
add_bullets(doc, [
    "Time expressions are converted to minutes since midnight, allowing values such as 3pm, 3:00 PM, and 15:00 to compare equally.",
    "Currency symbols, commas, and whitespace are removed before numeric comparison.",
    "Integer and floating-point representations are compared numerically.",
    "Lists are normalized item by item and sorted before comparison.",
    "Plain strings are trimmed and compared case-insensitively.",
])

doc.add_heading("Model and Configuration", level=1)
add_table(doc,
    ["Setting", "Value or source", "Behavior"],
    [
        ("Default model", "gemma-4-26b-a4b-it", "Used when no model override is supplied."),
        ("Model override", "INTENTLEDGER_MODEL", "Allows any configured model identifier without changing source code."),
        ("API credential", "GEMINI_API_KEY", "Read from the environment and never displayed by the application."),
        ("Generation", "temperature 0.0 and max 1024 tokens", "Encourages deterministic structured output."),
        ("Verification", "Model list request", "Separates key presence from successfully verified live model access."),
    ], [1.35, 2.6, 3.2]
)

doc.add_heading("Error Classification", level=1)
doc.add_paragraph(
    "The extractor preserves distinct operational failures so the UI can recommend an appropriate next action. Offline mode is returned when no API key is set. Other classifications include missing dependency, quota exhaustion, invalid API key, model not found, general model-call failure, and non-JSON output. The redesigned result panel shows the original error text and a targeted troubleshooting step without converting the error into a false evaluation score."
)

doc.add_heading("User Interface and Interaction Design", level=1)
doc.add_paragraph(
    "The interface uses a dark indigo developer-tool visual system with violet and electric-blue accents. Its layout prioritizes fast comprehension during a live demo while keeping detailed evidence available."
)
add_bullets(doc, [
    "Compact sidebar with branding, scenario navigation, model name, API-key presence, separate live-verification state, and benchmark limitations.",
    "Dashboard header with run-all actions and the actual configured model label.",
    "State-driven overview metrics for scenario count, completed scenarios, and average baseline and improved match rates; unrun evaluations display Not run rather than a misleading zero percent.",
    "Scenario workspace with transcript, expected state, baseline, improved, and compare tabs.",
    "Conversation timeline that visually distinguishes user and assistant turns.",
    "Field-level result table with consistent match, mismatch, missing, and unexpected styling.",
    "All-case matrix showing each scenario's baseline status, improved status, score, and measured change.",
    "Responsive layouts, keyboard-accessible native Streamlit controls, visible focus treatment, and session-state persistence across normal reruns.",
])

doc.add_heading("Session State and Interaction Behavior", level=1)
doc.add_paragraph(
    "Baseline and improved results are stored in separate dictionaries keyed by case identifier. Running a single case updates only that case and mode. Running all cases replaces the selected mode's result collection after processing each scenario with progress feedback. A scenario is considered completed when either mode has a stored result. The application also stores the API verification outcome and message independently from key presence."
)

doc.add_heading("Testing and Verification", level=1)
doc.add_paragraph(
    "The current automated suite contains 46 passing tests. It exercises the default model contract, environment override, API-key parsing, offline behavior, JSON fence removal, embedded JSON extraction, normalization rules, scoring outcomes, clarification handling, and aggregate summaries. No live API call is required for the tests."
)
add_table(doc,
    ["Test area", "Examples of verified behavior"],
    [
        ("Configuration", "Default Gemma identifier, arbitrary model override, absent and whitespace API keys."),
        ("Extraction helpers", "Plain, fenced, embedded, empty, and invalid JSON text."),
        ("Time normalization", "3pm equals 3:00 PM and 15:00; noon and midnight edge cases."),
        ("Numeric normalization", "Currency symbols, comma separators, integers, floats, and numeric strings."),
        ("Scoring", "Exact match, mismatch, missing field, extra field, ERROR result, and partial match rate."),
        ("Clarification", "Required clarification present versus missing."),
        ("Summary", "Empty result collection and mixed PASS, FAIL, and ERROR results."),
    ], [1.65, 5.5]
)
doc.add_paragraph(
    "The Streamlit application was also launched locally and inspected in a browser. Empty states, the offline ERROR path, dynamic model labeling, navigation, responsive styling, and removal of light-theme style leakage were checked visually. Live API success behavior remains dependent on a valid external credential and model access."
)

doc.add_heading("Security and Data Handling", level=1)
add_bullets(doc, [
    "The API key is read from an environment variable and is never rendered or logged by the application.",
    "Key presence is not presented as proof of connectivity; live verification has a separate status.",
    "Benchmark content is escaped before insertion into custom HTML rendering paths.",
    "The model receives only the selected benchmark transcript and prompt required for extraction.",
    "The project does not persist model results to a database; results remain in the Streamlit browser session state.",
])

doc.add_heading("Known Limitations", level=1)
add_numbered(doc, [
    "The benchmark contains only five hand-crafted English cases and cannot establish general model reliability.",
    "Prompt wording may influence outcomes, and the improved prompt is not proof of generalization.",
    "The model must return parseable JSON; the application does not use a strict structured-output schema.",
    "Cases are short and do not test long-context, multilingual, adversarial, or multi-agent conversations.",
    "Only one configured model is evaluated at a time.",
    "Live results depend on external API availability, account quota, credential validity, and model access.",
])

doc.add_heading("Recommended Demonstration Flow", level=1)
add_numbered(doc, [
    "Open the dashboard and explain the five-case scope and the distinction between key configured and live access verified.",
    "Select Accepted Correction and show the transcript and expected state before running either prompt.",
    "Run the baseline case, then the improved case, and inspect the field-level differences in Compare.",
    "Open Ambiguous Instruction to explain why a model guess is not equivalent to confirmed user intent.",
    "Run all cases for both modes when credentials and quota permit, then use the all-case matrix to summarize measured changes.",
    "Close by emphasizing that IntentLedger is an explainable demo benchmark, not a universal reliability score."
])

doc.add_heading("Future Enhancements", level=1)
add_bullets(doc, [
    "Expand the benchmark with larger, independently reviewed scenario sets and severity labels.",
    "Add structured-output schema support to reduce JSON parsing failures.",
    "Support side-by-side evaluation of multiple model configurations and repeated runs.",
    "Export results as JSON or CSV with prompt and model metadata for reproducibility.",
    "Add latency and token-usage measurements using real API response metadata.",
    "Introduce longer conversations, multilingual cases, and adversarial instruction conflicts.",
    "Add continuous-integration checks for the Streamlit presentation layer and accessibility behavior.",
])

doc.add_heading("Project Structure", level=1)
project_structure = (
    "IntentLedger-HackDay/\n"
    "  app.py                     Streamlit application flow and session state\n"
    "  ui.py                      Reusable renderers and visual design system\n"
    "  engine/\n"
    "    extractor.py             Model configuration, API call, and error handling\n"
    "    prompts.py               Baseline and improved prompt definitions\n"
    "    normalizer.py            Canonical value comparison\n"
    "    scorer.py                Field-level scoring and summaries\n"
    "    loader.py                Benchmark case loading\n"
    "  tests/\n"
    "    cases/                   Five scenario JSON files\n"
    "    test_extractor_config.py Configuration and parsing tests\n"
    "    test_normalizer.py       Normalization tests\n"
    "    test_scorer.py           Scoring and summary tests\n"
    "  requirements.txt          Python dependencies\n"
    "  README.md                 Setup, architecture, and usage guidance"
)
p = doc.add_paragraph()
run = p.add_run(project_structure)
run.font.name = "Consolas"
run.font.size = Pt(8.5)

doc.add_heading("Setup and Commands", level=1)
doc.add_paragraph("Install dependencies:")
p = doc.add_paragraph()
r = p.add_run("pip install -r requirements.txt")
r.font.name = "Consolas"
doc.add_paragraph("Set the API key in PowerShell when live evaluation is required:")
p = doc.add_paragraph()
r = p.add_run('$env:GEMINI_API_KEY = "your-api-key"')
r.font.name = "Consolas"
doc.add_paragraph("Optionally override the model:")
p = doc.add_paragraph()
r = p.add_run('$env:INTENTLEDGER_MODEL = "model-id"')
r.font.name = "Consolas"
doc.add_paragraph("Launch and test:")
p = doc.add_paragraph()
r = p.add_run("streamlit run app.py\npytest tests -v")
r.font.name = "Consolas"

doc.add_heading("Conclusion", level=1)
doc.add_paragraph(
    "IntentLedger demonstrates a practical way to make conversational intent tracking measurable and inspectable. Its value is not a single headline score; it is the trace from conversation to expected state, model prediction, normalized comparison, and verdict. The modular engine, explicit error handling, real-state metrics, and redesigned evaluation workspace provide a credible foundation for a hackathon demonstration and for future expansion into a broader evaluation suite."
)

# Footer with page field.
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IntentLedger Project Report   |   ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)

doc.core_properties.title = "IntentLedger Project Report"
doc.core_properties.subject = "Technical project report for the IntentLedger AI intent tracking benchmark"
doc.core_properties.keywords = "IntentLedger, Streamlit, AI evaluation, intent tracking, Gemma"
doc.save(OUT)
print(OUT)
